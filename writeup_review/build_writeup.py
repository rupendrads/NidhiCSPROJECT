"""Merge the student's edited write-up with Appendices E-H and fix the issues found."""
import copy, re, sys, zipfile
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

BASE = "Project_Writeup_BEFORE_EDIT.docx"
SRC = "Project_Writeup.docx"          # old scaffold that carries Appendices E-H
OUT = sys.argv[1] if len(sys.argv) > 1 else "Project_Writeup_MERGED.docx"
DASHBOARD = "frontend-dashboard.png"

doc = Document(BASE)
src = Document(SRC)
log = []


# ---------- helpers ----------
def find_para(pred, start=0):
    for i, p in enumerate(doc.paragraphs[start:], start):
        if pred(p.text):
            return i, p
    raise LookupError("paragraph not found")


def set_para_text(p, text):
    """Replace the paragraph's text, keeping the first run's formatting."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_para_lead(p, lead, rest):
    """First run keeps its (bold) formatting for the lead-in; the rest is added un-bold."""
    set_para_text(p, lead)
    r0 = p.runs[0]
    r1 = p.add_run(" " + rest)
    rpr = copy.deepcopy(r0._r.find(qn('w:rPr'))) if r0._r.find(qn('w:rPr')) is not None else None
    if rpr is not None:
        for b in rpr.findall(qn('w:b')) + rpr.findall(qn('w:bCs')):
            rpr.remove(b)
        r1._r.insert(0, rpr)


def replace_in_para(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    if old in p.text:                      # spans runs -> collapse
        set_para_text(p, p.text.replace(old, new))
        return True
    raise LookupError(f"{old!r} not in paragraph {p.text[:50]!r}")


def set_cell(cell, text, plain=False):
    p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    for extra in cell.paragraphs[1:]:
        extra._p.getparent().remove(extra._p)
    set_para_text(p, text)
    if plain:                                # drop inherited code/colour run formatting
        for r in p.runs:
            rpr = r._r.find(qn('w:rPr'))
            if rpr is not None:
                r._r.remove(rpr)


def insert_para_after(p, text, style=None):
    new = copy.deepcopy(p._p)
    for child in list(new):
        if child.tag != qn('w:pPr'):
            new.remove(child)
    p._p.addnext(new)
    np_ = Paragraph(new, p._parent)
    if style:
        np_.style = doc.styles[style]
    np_.add_run(text)
    return np_


def insert_picture_after(p, path, width_in):
    new = insert_para_after(p, "")
    new.runs[0].add_picture(path, width=Inches(width_in))
    return new


def numids_in(path):
    try:
        x = zipfile.ZipFile(path).read('word/numbering.xml').decode()
    except KeyError:
        return set()
    return set(re.findall(r'<w:num w:numId="(\d+)"', x))


# ---------- 1. truncated heading ----------
_, p = find_para(lambda t: t.startswith("1.2 Why the problem is amenable"))
set_para_text(p, "1.2 Why the problem is amenable to a computational solution")
log.append("§1.2 heading completed")

# ---------- 2. §1.6 login contradiction ----------
_, p = find_para(lambda t: t.startswith("Single-user, single-machine desktop tool."))
set_para_lead(p, "Single-machine desktop tool with a lightweight local login.", "Justified: it models one "
                 "investor (Krish) on his own laptop. A simple username/password (§2.3, §2.4) keeps his "
                 "data separate from anyone else who tries the tool, but there are no cloud accounts, "
                 "password-recovery flows or role permissions, which removes significant complexity.")
log.append("§1.6 reconciled with the login design in §2.3/§2.4")

# ---------- 3. 'Three tables' ----------
_, p = find_para(lambda t: t.startswith("Persistent storage - SQLite"))
replace_in_para(p, "Three tables:", "Four tables:")
log.append("§2.3 table count corrected")

# ---------- 4. §2.1.3 placeholder ----------
_, p = find_para(lambda t: "[YOUR INPUT: one or two sentences on which module" in t)
replace_in_para(p, p.text[p.text.index("[YOUR INPUT"):],
                "I will build persistence (database.py, sub-problem 1.1) first, because SC1 underpins "
                "everything else and it can be tested with no network at all; then the Zerodha layer, "
                "since every later stage needs real prices; then the risk maths; then the Module 1 front "
                "end; and finally the backtest engine, which consumes all of the above (this is the stage "
                "order followed in §2.7.1 and §3).")
log.append("§2.1.3 build-order placeholder filled")

# (step 5 removed: the §2.3.2 table already holds the student's content inside content controls)

# ---------- 6. API contract table: add the endpoints the code exposes ----------
_, h = find_para(lambda t: t.startswith("2.3.3 API contract"))
for el in h._p.itersiblings():
    if el.tag == qn('w:tbl'):
        api = [t for t in doc.tables if t._tbl is el][0]
        break
extra = [
    ("GET /api/status", "Health check", "-", "{connected, instrumentCount, lastError}"),
    ("GET /api/quote/<ticker>", "Latest traded price for one stock", "-", "{ticker, ltp}"),
    ("GET /api/portfolio/closed", "Realised (booked) positions", "-", "[ {ticker, quantity, avgBuyPrice, sellPrice, closeDate} ]"),
    ("GET /api/tickers", "Valid NSE symbols for the form", "-", "[ tradingsymbol, ... ]"),
    ("GET /api/history/<ticker>", "Daily closes (Module 2 data)", "?from&to", "[ {date, close} ]"),
]
template = api.rows[-1]._tr
for vals in extra:
    new_tr = copy.deepcopy(template)
    api._tbl.append(new_tr)
    row = api.rows[-1]
    for c, v in zip(row.cells, vals):
        set_cell(c, v)
log.append("§2.3.3 API table extended with 5 endpoints the backend exposes")

# ---------- 7. §2.5 fixes ----------
_, p = find_para(lambda t: t.startswith("(c) SMA-crossover backtest"))
replace_in_para(p, "(SC4,SC5,  SC7)", "(SC4, SC5, SC7)")
_, p = find_para(lambda t: t.startswith("Baseline initialization point."))
set_para_lead(p, "Baseline initialization point.", "The buy-and-hold benchmark buys on the first day of the "
                 "series (index 0, i.e. benchShares = floor(capital / close[0]) above, exactly as backtest.py "
                 "does). Justified: it measures what the investor would have got by buying at the start of the "
                 "window and doing nothing. Because the strategy cannot trade until both SMAs exist "
                 "(index ≥ longWindow), this is a deliberately conservative comparison that slightly favours "
                 "buy-and-hold - if the strategy still wins, the evidence is stronger.")
log.append("§2.5 benchmark start reconciled with pseudocode, §2.7.2 and backtest.py")

# flowchart caption
i, h = find_para(lambda t: t.startswith("2.5 Algorithms"))
for p in doc.paragraphs[i:]:
    if p._p.xpath('.//w:drawing'):
        insert_para_after(p, "Figure 2.1 - Flowchart of the /api/backtest route: validate parameters -> fetch "
                             "closes -> check history length -> short and long SMA -> crossover loop -> "
                             "buy-and-hold baseline -> stats -> verdict -> JSON reply. The two decision "
                             "diamonds are the error paths tested in §2.7.3 rows 10-11.", "Table Caption")
        break
log.append("§2.5 flowchart captioned")

# ---------- 8. §2.6 fixes ----------
_, p = find_para(lambda t: "[YOUR INPUT: quote your stakeholder's actual feedback.]" in t)
replace_in_para(p, "[YOUR INPUT: quote your stakeholder's actual feedback.]",
                "This follows Krish's own words in the interview (§1.3): \"please keep the screen simple and "
                "clean\" and, on market swings, \"it makes me anxious\".")
i, _ = find_para(lambda t: t.startswith("2.6 Usability"))
imgs = [p for p in doc.paragraphs[i:i + 40] if p._p.xpath('.//w:drawing')]
insert_para_after(imgs[0], "Figure 2.2 - Login screen wireframe. Two fields and one button only: username and "
                           "password rules from §2.4 are checked inline before the request is sent.", "Table Caption")
insert_para_after(imgs[1], "Figure 2.3 - Portfolio screen wireframe. Top-to-bottom order follows importance for "
                           "the stakeholder: total value (numbers are the hero), allocation ribbon, holdings beside "
                           "sector mix and return-by-holding, closed positions, and the collapsible strategy "
                           "backtest last (progressive disclosure).", "Table Caption")
for lbl, new in (("LOGIN SCREEN:", "Login screen:"), ("PORTFOLIO SCREEN", "Portfolio screen:")):
    _, p = find_para(lambda t, l=lbl: t.strip() == l)
    set_para_text(p, new)
log.append("§2.6 wireframes captioned, stakeholder-quote placeholder filled")

# remove the stale scaffold heading "2.6 Test strategy & test data" (lives in a content control)
for el in list(doc.element.body):
    if el.tag == qn('w:sdt') and "".join(el.itertext()).startswith("2.6 Test strategy & test data"):
        el.getparent().remove(el)
        log.append("stray '2.6 Test strategy & test data' heading removed")
        break

# move the misplaced testing paragraph under §2.7
_, stray = find_para(lambda t: t.startswith("Testing happens during development (iterative)"))
_, h27 = find_para(lambda t: t.startswith("2.7 Test strategy"))
h27._p.addnext(stray._p)
log.append("§2.6 stray 'Testing happens...' paragraph moved under §2.7")

# ---------- 9. §2.7.1 stage order aligned with §3 ----------
_, h = find_para(lambda t: t.startswith("2.7.1 Iterative"))
for el in h._p.itersiblings():
    if el.tag == qn('w:tbl'):
        it = [t for t in doc.tables if t._tbl is el][0]
        break
stages = [
    ("1", "database.py + /holdings routes (dummy prices)", "Add/merge/sell on a test DB; restart and re-read",
     "10 @100 then 5 @130", "qty 15, avg 110.00; row survives restart (SC1)"),
    ("2", "market_data.py (Zerodha)", "Mock API responses + one live call", "valid / invalid token",
     "prices returned / graceful error"),
    ("3", "analysis.py risk maths", "Unit tests against hand-calculated values", "[10,12,14,16,18]",
     "returns 0.2000, 0.1667, 0.1429, 0.1250; volatility = spreadsheet (SC3)"),
    ("4", "Module 1 front end (app.js)", "Manual + validation tests", "hostile input",
     "messages shown, no crash (SC8)"),
    ("5", "backtest.py SMA + crossover engine, backtest.js", "Unit test with the hand-computable series (§2.7.2)",
     "closes [10,9,8,9,11,12,10,8], 2/3, ₹1000", "SMA [None,None,12,14,16] for the 5-point series; 1 BUY @11, 1 SELL @8, final ₹730"),
    ("6", "Routes (app.py) end-to-end", "Browser tests incl. bad input", "see §2.7.3", "correct status codes"),
]
for r, vals in zip(it.rows[1:], stages):
    for c, v in zip(r.cells, vals):
        set_cell(c, v, plain=True)
log.append("§2.7.1 stage order now matches §3")

_, p = find_para(lambda t: "[Record user details and dates.]" in t)
replace_in_para(p, "[Record user details and dates.]",
                "The stakeholder session is recorded in §4.1 (T5) and §4.6.")

# ---------- 10. §2.8 traceability ----------
_, h = find_para(lambda t: t.startswith("2.8 Traceability"))
for el in h._p.itersiblings():
    if el.tag == qn('w:tbl'):
        tm = [t for t in doc.tables if t._tbl is el][0]
        break
trace = [
    ("Success criterion (§1.8)", "Addressed by"),
    ("SC1 Portfolio persists across a restart", "§2.3 (holdings table), §2.5(a), §2.5(f), §2.7.1 stage 1"),
    ("SC2 Value & return match a hand calculation", "§2.3.2 (Holding class), §2.5(a)"),
    ("SC3 Volatility & Sharpe match a spreadsheet", "§2.5(d), §2.7.1 stage 3"),
    ("SC4 Trades on the exact crossover days", "§2.5(b), §2.5(c), §2.7.2"),
    ("SC5 Labelled equity-vs-benchmark chart", "§2.5(c), §2.6"),
    ("SC6 Correct trade stats & drawdown", "§2.5(e), §2.5(g), §2.7.2"),
    ("SC7 Beat-the-benchmark verdict", "§2.5(c), §2.7.2"),
    ("SC8 Invalid input rejected clearly", "§2.4, §2.7.3"),
    ("NFR1 Accessibility / NFR2 Indian conventions", "§2.6"),
]
for r, vals in zip(tm.rows, trace):
    for c, v in zip(r.cells, vals):
        set_cell(c, v)
log.append("§2.8 traceability matrix: non-existent 2.3.1 removed, labels aligned to §1.8")

# ---------- 11. stale cross-references in §3 / §4 ----------
_, p = find_para(lambda t: "see 2.4d" in t)
replace_in_para(p, "see 2.4d", "see §2.5(d)")
_, p = find_para(lambda t: "(§2.4c)" in t)
replace_in_para(p, "(§2.4c)", "(§2.5(c))")
_, p = find_para(lambda t: "usability feature from §2.5" in t)
replace_in_para(p, "usability feature from §2.5", "usability feature from §2.6")
log.append("stale §2.4/§2.5 cross-references in §3.4, §3.6, §4.3 corrected")

# ---------- 12. dashboard screenshot at §3.5 ----------
_, p = find_para(lambda t: t.startswith("📸 [SCREENSHOT] The finished dashboard"))
set_para_text(p, "Figure 3.1 - The finished Module 1 dashboard (front end in mock mode with representative "
                 "sample data). Visible: hero value with 30-session sparkline; stat strip (invested, market "
                 "value, today's change, annualised volatility 16.8%, Sharpe 1.42); allocation ribbon; "
                 "holdings table; sector-mix donut; return-by-holding bars; closed-positions ledger; and the "
                 "collapsed strategy-backtest panel. Captured with prefers-reduced-motion so the count-up "
                 "shows its final figure.")
pic = insert_picture_after(p, DASHBOARD, 5.5)
p._p.addprevious(pic._p)           # picture first, caption under it
insert_para_after(p, "📸 [SCREENSHOT still needed] The Add form showing a validation error (SC8).")
log.append("§3.5 dashboard screenshot inserted with caption")

# ---------- 13. appendices B/C ----------
_, p = find_para(lambda t: t.startswith("B. Setup & run guide"))
set_para_text(p, "B. Setup & run guide — see Appendix E (Prerequisites & Setup) in this document.")
_, p = find_para(lambda t: t.startswith("C. Technical design overview"))
set_para_text(p, "C. Technical design overview — see Appendix F (Functionalities & API Reference) in this document.")

# ---------- 14. copy Appendices E-H ----------
have_nums = numids_in(BASE)
body = doc.element.body
sect = body.find(qn('w:sectPr'))
started = False
copied = 0
for el in list(src.element.body):
    if el.tag == qn('w:sectPr'):
        continue
    text = "".join(el.itertext())
    if not started and el.tag == qn('w:p') and text.startswith("Appendix E"):
        started = True
    if not started:
        continue
    new = copy.deepcopy(el)
    if new.xpath('.//w:drawing'):        # Appendix H picture: re-add against this package
        holder = Paragraph(new, doc._body)
        for child in list(new):
            if child.tag != qn('w:pPr'):
                new.remove(child)
        holder.add_run().add_picture(DASHBOARD, width=Inches(5.5))
    for bm in new.xpath('.//w:bookmarkStart|.//w:bookmarkEnd'):   # avoid id clashes with the base doc
        bm.set(qn('w:id'), str(9000 + int(bm.get(qn('w:id')))))
    for numpr in new.xpath('.//w:numPr'):
        nid = numpr.find(qn('w:numId'))
        if nid is not None and nid.get(qn('w:val')) not in have_nums:
            numpr.getparent().remove(numpr)
    sect.addprevious(new)
    copied += 1
log.append(f"Appendices E-H appended ({copied} blocks) and the Appendix H screenshot re-embedded")

# ---------- 15. remove the scaffold reminder line ----------
_, p = find_para(lambda t: t.startswith("Reminder: rewrite in your own words"))
p._p.getparent().remove(p._p)
log.append("scaffold 'Reminder' line removed from Appendices")

# ---------- 15c. login documentation ----------
def new_para_before(anchor_el, text, style_id):
    pr = anchor_el.makeelement(qn('w:p'), {})
    anchor_el.addprevious(pr)
    para = Paragraph(pr, doc._body)
    para.style = doc.styles[style_id]
    para.add_run(text)
    return para

def new_table_before(anchor_el, rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = doc.styles['Table']
    for ri, (r, vals) in enumerate(zip(tbl.rows, rows)):
        for c, v in zip(r.cells, vals):
            c.paragraphs[0].style = doc.styles['Compact']
            run = c.paragraphs[0].add_run(v)
            if ri == 0:
                run.bold = True
    anchor_el.addprevious(tbl._tbl)
    return tbl

LOGIN_IMG = "writeup_review/img/login-login.png"
PW_IMG = "writeup_review/img/login-pw.png"

def add_row_like(table, vals, before_last=False):
    tr = copy.deepcopy(table.rows[-2 if before_last else -1]._tr)
    if before_last:
        table.rows[-1]._tr.addprevious(tr)
    else:
        table._tbl.append(tr)
    row = table.rows[-2 if before_last else -1]
    for c, v in zip(row.cells, vals):
        set_cell(c, v, plain=True)

def table_after(heading_pred):
    _, h = find_para(heading_pred)
    for el in h._p.itersiblings():
        if el.tag == qn('w:tbl'):
            return [t for t in doc.tables if t._tbl is el][0]
    raise LookupError("table not found")

# §1.7 FR12
_, p = find_para(lambda t: t.startswith("FR11 - reject invalid input"))
insert_para_after(p, "FR12 - accounts: create an account, sign in, change password; each user sees only their own holdings")

# §2.1.2 tree
from docx.text.run import Run
_, p = find_para(lambda t: "1.6 Validate all inputs" in t and "Module 2" in t)   # the tree is one paragraph
runs = p.runs
for idx, r in enumerate(runs):
    if "1.6 Validate" in r.text:
        r.text = r.text.replace("└──", "├──")
        # the 1.6 line may be split over several runs: insert before the next line break
        j = idx + 1
        while j < len(runs) and runs[j]._r.find(qn('w:br')) is None:
            j += 1
        nr = copy.deepcopy(r._r)
        for child in list(nr):
            if child.tag != qn('w:rPr'):
                nr.remove(child)
        runs[j]._r.addprevious(nr)
        Run(nr, p).text = chr(10) + "│   └── 1.7 Accounts: register / sign in / change password       → database"
        break

# §2.2 file table
ft = table_after(lambda t: t.startswith("2.2 System architecture"))
add_row_like(ft, ("", "auth.js", "Sign-in / create-account screen, user chip + menu, change-password modal; keeps the token"))
add_row_like(ft, ("", "auth.py", "Password hashing, signed login tokens, require_login route guard"))

# §2.3.3 API contract
_, h = find_para(lambda t: t.startswith("2.3.3 API contract"))
insert_para_after(h, "Authentication: a successful login returns a signed, time-limited token; the browser sends it as "
                     "Authorization: Bearer <token> on every other request and the server rejects anything else with 401. "
                     "A token rather than a session cookie because the page (port 8000) and the API (port 5000) are "
                     "different origins.")
api = table_after(lambda t: t.startswith("2.3.3 API contract"))
for r in api.rows:
    if r.cells[0].text.strip().startswith("POST /api/login"):
        for c, v in zip(r.cells, ("POST /api/auth/login", "Sign in", "{username, password}", "200 + {token, user}")):
            set_cell(c, v)
add_row_like(api, ("POST /api/auth/register", "Create account", "{username, password}", "201 + {token, user}"))
add_row_like(api, ("GET /api/auth/me", "Who am I (token check)", "-", "{id, username, createdAt}"))
add_row_like(api, ("POST /api/auth/change-password", "Change password", "{currentPassword, newPassword}", "200 {ok}"))

# §3.7 table: two more rows before the "…" row
tt = table_after(lambda t: t.startswith("3.7 Testing to inform development"))
add_row_like(tt, ("6", "Login guard", "no token; forged token; wrong password", "401", "401", "✓", "—"), before_last=True)
add_row_like(tt, ("6", "Session cookie across ports", "cookie from :5000, page on :8000", "sent", "not sent", "✗",
                  "switched to a signed bearer token in the Authorization header → fixed"), before_last=True)

# §3.8 Stage 6 (before the challenges section, which becomes 3.9)
_, h4 = find_para(lambda t: t.strip() == "4. Evaluation")   # challenges (3.9) are inserted before §4 later, i.e. after this
A = h4._p
new_para_before(A, "3.8 Stage 6 — Accounts, sign-in and password change (FR12)", "Heading2")
new_para_before(A, "Sub-problem 1.7. Built last, once every other module was proven, because it wraps them rather than "
                   "feeding them. Back end: a users table (username UNIQUE, salted PBKDF2 password_hash, created_at), "
                   "four routes under /api/auth (register, login, me, change-password) and a require_login decorator "
                   "that every data route now carries; holdings and closed_positions gained a user_id so each account "
                   "sees only its own rows. Front end: auth.js draws a full-screen sign-in / create-account card with "
                   "floating labels, a password-strength meter and a show/hide toggle, then hands over to the dashboard "
                   "with a short zoom-out; a user chip in the top bar opens Change password and Sign out.", "BodyText")
new_para_before(A, "Annotated example (the three ideas that make it safe: hash, sign, guard):", "BodyText")
for line in [
    "def hash_password(password):            # salted PBKDF2 - the password itself is never stored",
    "    return generate_password_hash(password)",
    "",
    "def issue_token(user_id):               # signed + time-limited; forged or expired -> None",
    "    return _serializer.dumps({\"uid\": user_id})",
    "",
    "def require_login(view):                # 401 unless Authorization: Bearer <valid token>",
    "    @wraps(view)",
    "    def wrapper(*a, **kw):",
    "        uid = current_user_id()",
    "        if uid is None:",
    "            return jsonify({\"error\": \"Please sign in.\"}), 401",
    "        g.user_id = uid                 # every db call is scoped to this",
    "        return view(*a, **kw)",
    "    return wrapper",
]:
    new_para_before(A, line, "SourceCode")
pic = new_para_before(A, "", "BodyText"); pic.add_run().add_picture(LOGIN_IMG, width=Inches(6))
new_para_before(A, "Figure 3.2 - The sign-in screen. Three slow-drifting colour fields and a gold equity curve that "
                   "draws itself on load (all disabled under prefers-reduced-motion); the tab ink slides between "
                   "Sign in and Create account; labels float up on focus; wrong details shake the card and show a "
                   "clear message.", "Table Caption")
pic = new_para_before(A, "", "BodyText"); pic.add_run().add_picture(PW_IMG, width=Inches(6))
new_para_before(A, "Figure 3.3 - Change password from the user-chip menu: current password is re-checked on the "
                   "server, the new one must be at least 8 characters and different, and a strength bar gives "
                   "feedback while typing; success draws an animated tick.", "Table Caption")
new_para_before(A, "Test: no token, a forged token and a wrong password all return 401 / 'Wrong username or "
                   "password' (the same message for unknown user and wrong password, so nothing is revealed); a second "
                   "account sees an empty portfolio (row isolation); after a password change the old password is "
                   "rejected and the new one accepted (see §3.7 rows 6).", "BodyText")
new_para_before(A, "Review: FR12 complete; the design in §2.3 / §2.4 / §2.6 now matches the build.", "BodyText")
log.append("§3.8 Stage 6 (accounts) inserted; challenges renumbered to §3.9")

# Appendix E.3
_, p = find_para(lambda t: t.startswith("2.  Open http://localhost:8000"))
_, p3 = find_para(lambda t: t.startswith("3.  To stop:"))
set_para_text(p3, "4.  To stop: press Ctrl+C in the backend window and close the front-end window.")
insert_para_after(p, "3.  First run: choose Create account on the sign-in screen (username 3-30 letters/digits/underscore, "
                     "password at least 8 characters). Any holdings saved before accounts existed are attached to this "
                     "first account. Sign in stays valid for 7 days; use the user chip (top right) to change your "
                     "password or sign out.")

# Appendix F.1 bullet + F.5 rows
_, p = find_para(lambda t: t.startswith("•  Indian conventions & validation"))
insert_para_after(p, "•  Accounts — create account, sign in (signed 7-day token), change password, sign out; every "
                     "holding and closed position belongs to the signed-in user.")
_, h = find_para(lambda t: t.startswith("F.5 REST API endpoints"))
insert_para_after(h, "All endpoints except /api/status, /api/tickers and /api/auth/* require the header "
                     "Authorization: Bearer <token> and answer 401 without it.")
f5 = table_after(lambda t: t.startswith("F.5 REST API endpoints"))
for vals in (("POST", "/api/auth/register", "Create account → {token, user}"),
             ("POST", "/api/auth/login", "Sign in → {token, user}"),
             ("GET", "/api/auth/me", "Current user (checks the token)"),
             ("POST", "/api/auth/change-password", "Change password (needs current password)")):
    add_row_like(f5, vals)

# Appendix G row
g1 = table_after(lambda t: t.startswith("G.1 Troubleshooting"))
add_row_like(g1, ("Every request answers 401 “Please sign in”",
                  "No token, the token expired (7 days), or backend/.secret_key was deleted/regenerated so old tokens no longer verify.",
                  "Sign in again (the screen reappears by itself); the user chip → Sign out clears a stale token."))
log.append("Appendices E/F/G updated for accounts")

# ---------- 15a. §3.8 challenges + Appendix I variables ----------
_, h4 = find_para(lambda t: t.strip() == "4. Evaluation")
A = h4._p
new_para_before(A, "3.9 Challenges met during development and how they were resolved", "Heading2")
new_para_before(A, "Each row is a problem that actually came up while building the system, the cause once it was "
                   "found, and the fix that is now in the code (file and function named so it can be checked). "
                   "The first row is the failed test already described in §3.3; the rest were found the same "
                   "way - by a stage not behaving as its test expected - and each is a justified remedial action.",
                "BodyText")
challenges = [
 ("#", "Challenge (symptom)", "Cause", "Resolution (where in the code)"),
 ("1", "Historical price requests returned no candles (Stage 2).",
  "The supplied helper Zerodha.get_instrument_tokens() maps a symbol to exchange_token, but Kite's historical API needs instrument_token.",
  "Bypassed the helper: market_data.fetch_nse_equity() keeps instrument_token from kite.instruments('NSE'); seed_instruments.py stores it in the instruments table; db.get_instrument_token() feeds historical_daily()."),
 ("2", "Login worked once, then every call failed the next day; each login needs a 2-factor code.",
  "Zerodha invalidates the access token every morning and requires TOTP on login.",
  "Zerodha.login() reuses the token in zerodhaAccessToken.json only if its saved date is today (load_saved_token), otherwise generates the code with pyotp.TOTP(secret).now(), logs in and caches the new token (save_token). MarketData._ensure() retries login once before any call."),
 ("3", "The app was unusable whenever Zerodha was unavailable (paid API, closed market, no internet).",
  "Every route depended on a live connection.",
  "'Degraded mode': every MarketData method fails softly and records last_error; get_holdings values a holding at avg_buy_price when there is no quote; portfolio_summary returns 0 risk; _load_prices() falls back to sample_data/<TICKER>.csv; the front end also has useMock for a no-backend demo."),
 ("4", "Login failed with 'Configuration file not found' depending on the folder the server was started from.",
  "The Zerodha class uses relative paths ('zerodhaConfig.json').",
  "config.py derives BACKEND_DIR, ZERODHA_DIR, SAMPLE_DIR and DB_PATH from __file__; MarketData.connect() sets z.config_path and z.token_path to absolute paths before login."),
 ("5", "/api/portfolio/summary took several seconds.",
  "Volatility and Sharpe need ~400 days of history for every holding - one API call per stock per page load.",
  "A module-level _history_cache {ticker: {date: close}} in app.py stores each series for the life of the process (_daily_closes())."),
 ("6", "Dates from the two price sources did not match, breaking sorting and the equity-curve x-axis.",
  "Zerodha candles carry datetime objects; CSV rows carry strings.",
  "_fmt_date() in app.py normalises both to YYYY-MM-DD before the data reaches backtest.py, which therefore does not care where prices came from."),
 ("7", "Portfolio volatility was distorted when holdings had different history lengths.",
  "Summing values on dates that only some stocks traded on produces jumps.",
  "portfolio_summary() takes the intersection of every holding's date set, requires more than two common dates, builds value_series = sum(quantity x close) per common date and passes it to analysis.portfolio_risk()."),
 ("8", "On a fresh install every ticker was rejected as 'not a valid NSE stock'.",
  "SC8 validation checks the instruments table, which starts empty.",
  "_validate_holding() only enforces the check when db.instruments_count() > 0; /api/tickers serves the built-in SECTORS list until the table is seeded."),
 ("9", "The symbol list was far larger than the ~1,800 real equities.",
  "kite.instruments('NSE') returns bonds, ETFs and warrants too.",
  "fetch_nse_equity() keeps only instrument_type == 'EQ' and segment == 'NSE'; the remaining noise is recorded as a maintenance item in §4.4."),
 ("10", "Front end (built first, on mock data) and back end used different field names.",
  "The UI expects camelCase JSON; SQLite columns are snake_case.",
  "config.js declares the REST contract up front and app.py matches it; _holding_json() / _closed_json() convert rows to camelCase, so going live was one flag: useMock: false."),
 ("11", "Every browser request failed although the API answered in a direct call.",
  "The page is served on port 8000 and the API on 5000 - blocked by the same-origin policy.",
  "flask_cors.CORS(app) in app.py adds the Access-Control-Allow-Origin headers."),
 ("12", "The backtest could act on the first day both SMAs existed, treat equal SMAs as a cross, and 'buy' zero shares with tiny capital.",
  "A crossover needs yesterday's SMAs, which are None early on; equality is not a cross; int(cash // price) can be 0.",
  "run_sma_backtest() only tests a cross when i > 0 and all four SMA values exist; comparisons are strict; the BUY is guarded by if shares > 0. simple_moving_average() uses a running sum so it is O(n)."),
 ("13", "No sector information for the sector-mix donut.",
  "Kite provides no sector field.",
  "sectors.py holds a lookup table with sector_of() defaulting to 'Other'; documented in §4.4 and §4.5."),
 ("14", "Start-up (database init + Zerodha login) ran twice.",
  "Flask's debug reloader imports the module twice.",
  "app.run(debug=True, use_reloader=False) keeps tracebacks but runs startup() once."),
 ("15", "A login session cookie set by the API (port 5000) was not sent back by the page (port 8000).",
  "The page and the API are different origins, so a cookie needs SameSite / CORS-credential exceptions.",
  "auth.py issues a signed, time-limited token (itsdangerous) instead; auth.js stores it and sends Authorization: Bearer <token> on every request; require_login() verifies it. Plain CORS is enough."),
 ("16", "Adding accounts would have orphaned holdings saved before accounts existed.",
  "The old holdings / closed_positions tables had no user_id column.",
  "database._migrate() adds the column on start-up and create_user() hands rows with user_id IS NULL to the first account that registers - nothing is lost when upgrading portfolio.db."),
]
new_table_before(A, challenges)
new_para_before(A, "Review: every challenge above was resolved inside the module responsible for it (§2.2), which is "
                   "why none of the fixes required changes in more than one file - evidence that the modular design worked.",
                "BodyText")
log.append("§3.9 challenges section (16 rows) inserted before §4")

sect = doc.element.body.find(qn('w:sectPr'))
new_para_before(sect, "Appendix I — Key Variables and Their Usage", "Heading1")
new_para_before(sect, "The variables that carry the state of the system, grouped by file. Names are exactly as they "
                      "appear in the source so the code can be cross-read with §2.3 and §2.5.", "BodyText")
groups = [
 ("I.1 backend/config.py — constants", [
   ("Variable", "Value", "Used for"),
   ("BACKEND_DIR / PROJECT_DIR / ZERODHA_DIR / SAMPLE_DIR", "derived from __file__", "absolute folders so the server runs from any working directory"),
   ("DB_PATH", "backend/portfolio.db", "the SQLite file opened by database.get_connection()"),
   ("EXCHANGE / CURRENCY", "'NSE' / 'INR'", "market identity; mirrored in frontend/config.js"),
   ("RISK_FREE_RATE", "6.5 (% p.a., Indian 10-year G-Sec)", "numerator of the Sharpe ratio"),
   ("TRADING_DAYS", "252", "annualising volatility (sqrt 252) and return (x 252)"),
   ("HOST / PORT", "localhost / 5000", "where Flask listens; must match config.js server"),
 ]),
 ("I.2 backend/market_data.py — MarketData", [
   ("Variable", "Type", "Used for"),
   ("market", "MarketData (single shared instance)", "the one Zerodha wrapper every route uses"),
   ("self.z", "Zerodha", "the underlying Kite session (self.z.kite)"),
   ("self.connected", "bool", "whether login succeeded; drives degraded mode and /api/status"),
   ("self.last_error", "str or None", "last failure message, shown by /api/status for debugging"),
   ("_ZERODHA_IMPORT_OK", "bool", "whether the Zerodha class / kiteconnect library could be imported"),
 ]),
 ("I.3 backend/app.py — routes", [
   ("Variable", "Type", "Used for"),
   ("_history_cache", "dict {ticker: {date: close}}", "in-memory cache so history is fetched once per stock per run"),
   ("prices", "dict {ticker: ltp}", "batched live prices from _prices_for_tickers()"),
   ("ltp", "float", "price used to value a holding; falls back to avg_buy_price"),
   ("quotes", "dict {ticker: {last_price, prev_close, net_change}}", "today's change"),
   ("total_mv / total_inv / day_change / prev_value", "float accumulators", "portfolio totals in portfolio_summary()"),
   ("closes_by_ticker / common_dates / value_series", "dict / set / list", "building the portfolio value series for volatility and Sharpe"),
   ("token", "int or None", "instrument_token from the database, needed by the history API"),
   ("source", "'zerodha' | 'csv' | 'none'", "tells the UI where backtest prices came from"),
   ("short_w / long_w / capital", "int / int / float", "validated backtest parameters (defaults 20 / 50 / 100000)"),
 ]),
 ("I.4 backend/backtest.py — run_sma_backtest()", [
   ("Variable", "Type", "Used for"),
   ("dates / closes", "list[str] / list[float]", "parallel lists, oldest first; index i is one trading day"),
   ("short_sma / long_sma", "list[float or None]", "rolling means; None where history is too short"),
   ("running", "float", "running-sum accumulator in simple_moving_average() (O(n))"),
   ("cash", "float", "uninvested money; starts at capital"),
   ("shares", "int", "whole shares currently held: int(cash // price)"),
   ("position", "0 or 1", "0 = in cash, 1 = holding; prevents double buys / sells"),
   ("buy_price", "float", "entry price of the open lot, for the SELL profit"),
   ("crossed_up / crossed_down", "bool", "the strict-inequality crossover tests (SC4)"),
   ("trades", "list[dict]", "the trade log handed to Module 3 (analysis.trade_stats)"),
   ("equity_curve / benchmark_curve", "list[{date, value}]", "strategy and buy-and-hold value per day (SC5)"),
   ("bench_shares / bench_cash", "int / float", "buy-and-hold lot bought at closes[0]"),
   ("final_value / benchmark_final / beat", "float / float / bool", "the SC7 verdict"),
 ]),
 ("I.5 backend/analysis.py — risk maths", [
   ("Variable", "Type", "Used for"),
   ("rets", "list[float]", "daily returns (today - yesterday) / yesterday"),
   ("daily_sd", "float", "sample standard deviation of returns -> volatility"),
   ("mean_daily / annual_return / ann_vol", "float", "inputs to the Sharpe ratio"),
   ("peak / worst / drop", "float", "running peak and largest peak-to-trough fall -> max drawdown"),
   ("completed / wins", "list[dict]", "SELL trades and the profitable subset -> trade count, win rate, average profit"),
 ]),
 ("I.6 zerodha/Zerodha.py — session", [
   ("Variable", "Type", "Used for"),
   ("config_path / token_path", "str", "credentials file and cached-token file (set to absolute paths by market_data)"),
   ("credentials", "dict", "api_key, api_secret, user_id, password, totp secret"),
   ("kite", "KiteConnect", "the API client used for ltp / quote / historical_data / instruments"),
   ("saved_token / request_token / access_token", "str", "the login handshake; access_token is cached with today's date"),
 ]),
 ("I.7 frontend/config.js, app.js, backtest.js", [
   ("Variable", "Type", "Used for"),
   ("APP_CONFIG.useMock", "bool", "true = built-in demo data, no backend; false = live API"),
   ("APP_CONFIG.server / baseUrl", "object / getter", "where the API lives (must match config.py HOST / PORT)"),
   ("APP_CONFIG.endpoints", "object", "the REST contract both sides agree on"),
   ("APP_CONFIG.market.locale / symbol", "'en-IN' / rupee sign", "lakh / crore grouping and currency sign (NFR2)"),
   ("NSE_UNIVERSE", "object {ticker: {name, sector, ltp}}", "every known symbol with its live price; feeds quote(), sectorOf() and the dropdown"),
   ("state.holdings / state.closed / state.summary", "arrays / object", "the single UI state everything renders from"),
   ("_tickerSorted", "array or null", "cached sorted symbol list; set to null to invalidate after a reload"),
   ("_heroAnimated", "bool", "run the count-up once (skipped under prefers-reduced-motion)"),
   ("tickersLoaded", "bool", "backtest form loads the symbol list only once"),
   ("payload {ticker, shortWindow, longWindow, capital}", "object", "the JSON body sent to POST /api/backtest"),
 ]),
 ("I.8 backend/auth.py, frontend/auth.js — accounts", [
   ("Variable", "Type", "Used for"),
   ("SECRET_PATH / _serializer", "str / URLSafeTimedSerializer", "random signing secret in backend/.secret_key (git-ignored) and the token signer built from it"),
   ("TOKEN_MAX_AGE", "int (7 days, in seconds)", "how long a login token stays valid"),
   ("USERNAME_RE / MIN_PASSWORD", "regex / int", "the §2.4 rules: 3-30 letters/digits/underscore; at least 8 characters"),
   ("g.user_id", "int (Flask request context)", "set by require_login(); every database call is scoped to it"),
   ("password_hash", "str column in users", "pbkdf2:sha256 salted hash - the password itself is never stored"),
   ("AUTH (window.AUTH)", "object", "token(), user(), headers(), signedIn(), signOut(), expired() - shared by app.js and backtest.js"),
   ("pt.token / pt.user", "localStorage keys", "the bearer token and the signed-in user shown in the chip"),
   ("mode", "'login' | 'register'", "which tab of the sign-in card is active"),
 ]),
]
for title, rows in groups:
    new_para_before(sect, title, "Heading2")
    new_table_before(sect, rows)
log.append("Appendix I (key variables, 8 tables) appended")

# ---------- 15b. keep title page and table cells out of the TOC ----------
first = list(doc.element.body)[:4]
for el in first:
    for ps in el.iter(qn('w:pStyle')):
        if ps.get(qn('w:val')) == 'Heading1':
            ps.set(qn('w:val'), 'Title' if el.tag == qn('w:p') else 'Subtitle')
_, h = find_para(lambda t: t.startswith("2.3.2 In-program data structures"))
for el in h._p.itersiblings():
    if el.tag == qn('w:tbl'):
        for ps in el.iter(qn('w:pStyle')):
            if ps.get(qn('w:val')) == 'Heading2':
                ps.set(qn('w:val'), 'Compact')
        for r in el.find(qn('w:tr')).iter(qn('w:r')):        # bold header row
            rpr = r.find(qn('w:rPr'))
            if rpr is None:
                rpr = r.makeelement(qn('w:rPr'), {}); r.insert(0, rpr)
            if rpr.find(qn('w:b')) is None:
                rpr.insert(0, rpr.makeelement(qn('w:b'), {}))
        break
# prose / pseudocode blocks that were given a Heading style (they showed up in the TOC)
def restyle(el, style_id):
    for x in el.iter(qn('w:pStyle')):
        x.set(qn('w:val'), style_id)

def block_text(el):
    return "".join(el.itertext()).strip()

_, pa = find_para(lambda t: t.startswith("(a) Weighted-average lot merge"))
algo_style = pa.style.style_id
PROSE = ("Justification: encapsulating", "Standardised error responses", "Justification for architectural separation")
body_els = list(doc.element.body)
in_algos = False
demoted = 0
for el in body_els:
    if el.tag not in (qn('w:p'), qn('w:sdt')):
        continue
    t = block_text(el)
    if t.startswith("(f) Sell position"):
        in_algos = True
    if t.startswith("2.6 Usability"):
        in_algos = False
    ps = el.find('.//' + qn('w:pStyle'))
    if ps is None or not ps.get(qn('w:val')).startswith('Heading'):
        continue
    if t.startswith(PROSE):
        restyle(el, 'BodyText'); demoted += 1
    elif in_algos:
        if t.startswith(("(f)", "(g)", "(h)")):
            restyle(el, algo_style)
            for r in el.iter(qn('w:r')):
                rpr = r.find(qn('w:rPr'))
                if rpr is None:
                    rpr = r.makeelement(qn('w:rPr'), {}); r.insert(0, rpr)
                if rpr.find(qn('w:b')) is None:
                    rpr.insert(0, rpr.makeelement(qn('w:b'), {}))
        elif t.startswith("(") or el.find('.//' + qn('w:drawing')) is not None:
            restyle(el, 'BodyText')
        else:
            restyle(el, 'SourceCode')
        demoted += 1
log.append(f"{demoted} prose/pseudocode blocks moved off Heading styles (they were polluting the TOC)")
log.append("title page uses Title/Subtitle and §2.3.2 cells use body style, so the TOC lists headings only")

# ---------- 16. make every bookmark id unique (base file already had clashes) ----------
counter = [0]; stack = {}
for el in doc.element.body.iter():
    if el.tag == qn('w:bookmarkStart'):
        counter[0] += 1
        stack.setdefault(el.get(qn('w:id')), []).append(str(counter[0]))
        el.set(qn('w:id'), str(counter[0]))
    elif el.tag == qn('w:bookmarkEnd'):
        old_id = el.get(qn('w:id'))
        if stack.get(old_id):
            el.set(qn('w:id'), stack[old_id].pop(0))
log.append("bookmark ids renumbered uniquely")

doc.save(OUT)
print("\n".join("  - " + l for l in log))
print("saved", OUT)
